"""Build and audit the corrected 2026-09-27 HAX research-report bundle."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import shutil
import tempfile
from zipfile import ZIP_DEFLATED, ZipFile


PROJECT_ROOT = Path(__file__).resolve().parents[2]
BUNDLE = PROJECT_ROOT / "reports/hax_research_2026_09_27_v2"


def _paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def _replace_text(document, replacements: dict[str, str]) -> None:
    remaining = dict(replacements)
    for paragraph in _paragraphs(document):
        for old, new in tuple(remaining.items()):
            if old in paragraph.text:
                paragraph.text = paragraph.text.replace(old, new)
                remaining.pop(old)
    if remaining:
        raise ValueError(f"DOCX replacement text was not found: {tuple(remaining)}")


def _replace_media(docx_path: Path, media_name: str, image_path: Path) -> None:
    replacement = image_path.read_bytes()
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        found = False
        with ZipFile(docx_path, "r") as source, ZipFile(
            temporary, "w", ZIP_DEFLATED
        ) as target:
            for info in source.infolist():
                value = source.read(info.filename)
                if info.filename == media_name:
                    value = replacement
                    found = True
                target.writestr(info, value)
        if not found:
            raise ValueError(f"DOCX media target is missing: {media_name}")
        temporary.replace(docx_path)
    finally:
        temporary.unlink(missing_ok=True)


def patch_docx_bundle(bundle: Path = BUNDLE) -> None:
    """Apply the audited wording and regenerated figures to copied v1 documents."""

    from docx import Document

    selection = bundle / "thermotwin_selection_research_report_2026_09_27_v2.docx"
    shutil.copy2(
        PROJECT_ROOT
        / "reports/hax_research_2026_09_17_v1"
        / "thermotwin_selection_research_report_2026_09_17_v1.docx",
        selection,
    )
    document = Document(selection)
    _replace_text(
        document,
        {
            "September 17 2026": "September 27 2026",
            "approximately 4-7 times smaller percentage errors.": (
                "parameter-specific percentage errors smaller by factors of 6.55, "
                "6.60, and 3.92 (a 3.9-6.6-fold range)."
            ),
            "Repository files inspected 17 September 2026.": (
                "Repository files inspected 27 September 2026."
            ),
        },
    )
    document.save(selection)
    _replace_media(
        selection,
        "word/media/image10.png",
        bundle / "assets/selection_research_summary_2026_09_27.png",
    )

    pinn = bundle / "thermotwin_pinn_research_report_2026_09_27_v2.docx"
    shutil.copy2(
        PROJECT_ROOT
        / "reports/hax_research_2026_09_17_v1"
        / "thermotwin_pinn_research_report_2026_09_17_v1.docx",
        pinn,
    )
    document = Document(pinn)
    _replace_text(
        document,
        {
            "September 17 2026": "September 27 2026",
            (
                "Across ten paired trials, mean resistivity error is 1.7537% for "
                "the PINN and 4.6639% for the conventional estimator; the PINN "
                "performs better in nine trials. Predictions under a withheld "
                "operating condition are more evenly matched."
            ): (
                "Across the same ten paired trials, mean resistivity error is "
                "1.7537% for the PINN and 4.6639% for the conventional estimator; "
                "the PINN performs better in nine trials. The metric spans 285-315 "
                "K, while fitted trajectories reached only approximately "
                "294.83-305.21 K, leaving roughly two-thirds of the evaluation "
                "interval outside the training-temperature range. It therefore "
                "does not establish unique recovery of the whole curve. Predictions "
                "under a withheld operating condition are more evenly matched."
            ),
            (
                "Property error is relative RMSE at 61 equally spaced temperatures "
                "from 285 to 315 K. This is the declared evaluation interval, not "
                "the range fully excited by fitting. A fresh truth-only simulation "
                "during report preparation established a training temperature range "
                "of approximately 294.83-305.21 K. This numerical audit did not "
                "retrain or alter any estimator."
            ): (
                "Property error is relative RMSE at 61 equally spaced temperatures "
                "from 285 to 315 K. Fitted trajectories reached approximately "
                "294.83-305.21 K, leaving about 19.62 K of the 30 K evaluation "
                "interval—roughly two-thirds—outside the training-temperature "
                "range. This is extrapolative coverage, not two-thirds of the error "
                "magnitude, and it does not establish unique identification of the "
                "entire curve. A fresh truth-only simulation during report "
                "preparation established the range without retraining or altering "
                "any estimator."
            ),
            (
                "These are paired ten-trial means; substituting conventional "
                "twenty-trial averages would change the comparison's data composition."
            ): (
                "These are paired ten-trial means. The separate regularized "
                "conventional mean across all 20 available conventional trials is "
                "5.0747%; it is not substituted into the paired comparison."
            ),
            "Resistivity relative RMSE, 285-315 K": (
                "Resistivity relative RMSE, same ten paired trials, 285-315 K"
            ),
        },
    )
    document.save(pinn)
    _replace_media(
        pinn,
        "word/media/image6.png",
        bundle / "assets/pinn_research_summary_2026_09_27.png",
    )


def _file_record(path: Path) -> dict:
    value = path.read_bytes()
    return {
        "path": path.relative_to(PROJECT_ROOT).as_posix(),
        "bytes": len(value),
        "sha256": hashlib.sha256(value).hexdigest(),
    }


def write_manifest(bundle: Path = BUNDLE) -> Path:
    """Bind unchanged scientific sources and every revised bundle output."""

    v1_manifest = json.loads(
        (PROJECT_ROOT / "reports/hax_research_2026_09_17_v1/source_manifest.json")
        .read_text(encoding="utf-8")
    )
    outputs = sorted(
        path
        for path in bundle.rglob("*")
        if path.is_file() and path.name != "source_manifest.json"
    )
    manifest = {
        "schema_version": 2,
        "report_date": "2026-09-27",
        "bundle_version": "hax_research_2026_09_27_v2",
        "historical_bundle_preserved": "reports/hax_research_2026_09_17_v1",
        "scientific_inputs_changed": False,
        "scientific_input_note": (
            "The v1 scientific source records are retained verbatim below. Report "
            "wording, summary figures, DOCX, and PDF files were regenerated; no fit, "
            "training run, or scientific partition was rerun."
        ),
        "unchanged_scientific_sources_from_v1": v1_manifest["sources"],
        "revised_outputs": [_file_record(path) for path in outputs],
    }
    target = bundle / "source_manifest.json"
    target.write_text(
        json.dumps(manifest, sort_keys=True, indent=2) + "\n", encoding="utf-8"
    )
    return target


def main(argv=None) -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--patch-docx", action="store_true")
    parser.add_argument("--write-manifest", action="store_true")
    args = parser.parse_args(argv)
    if not args.patch_docx and not args.write_manifest:
        parser.error("choose --patch-docx or --write-manifest")
    if args.patch_docx:
        patch_docx_bundle()
    if args.write_manifest:
        print(write_manifest())


if __name__ == "__main__":
    main()
